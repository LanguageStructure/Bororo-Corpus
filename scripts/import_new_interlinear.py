#!/usr/bin/env python3
from pathlib import Path
from docx import Document
import csv,sys

FIELDS=['id','source','segmentation','gloss','portuguese','reviewed','segmentation_reviewed','gloss_reviewed','portuguese_reviewed','editorial_note']

def write(path,rows):
 path.parent.mkdir(parents=True,exist_ok=True)
 with path.open('w',encoding='utf-8',newline='') as f:
  w=csv.DictWriter(f,fieldnames=FIELDS,delimiter='\t',lineterminator='\n');w.writeheader();w.writerows(rows)

def groups(doc,skip=1):
 out=[];g=[]
 for p in doc.paragraphs[skip:]:
  t=p.text.strip()
  if t:g.append(t)
  elif g:out.append(g);g=[]
 if g:out.append(g)
 return out

def bible(path):
 rows=[];pending=''
 for g in groups(Document(path)):
  if len(g)>=4:
   rows.append(dict.fromkeys(FIELDS,'')|{'id':f'BOR-CORBO-BIBOR-{len(rows)+1:03d}','source':g[0],'segmentation':g[1],'gloss':g[2],'portuguese':g[3]})
  elif len(g)==3 and g[0].startswith('Amode '):
   rows.append(dict.fromkeys(FIELDS,'')|{'id':f'BOR-CORBO-BIBOR-{len(rows)+1:03d}','source':g[0],'segmentation':g[1],'portuguese':g[2]})
  elif len(g)==1 and rows: rows[-1]['editorial_note']=g[0]
 return rows

def paulinho(path):
 d=Document(path);p={i:x.text.strip() for i,x in enumerate(d.paragraphs) if x.text.strip()}
 starts=[2,8,14,19,24,30,36,40,45,51,57,63,69,77,82,88,93,98,101,105,108,112,116,121,126,134,140,147,153,165]
 ignore={0,12};rows=[]
 for n,st in enumerate(starts):
  end=starts[n+1] if n+1<len(starts) else 10**9
  vals=[p[i] for i in sorted(p) if st<=i<end and i not in ignore]
  src=vals.pop(0);seg=gl=pt=''
  if vals and ('=' in vals[0] or '∅' in vals[0] or '–' in vals[0]):seg=vals.pop(0)
  if len(vals)>=2:gl=vals.pop(0);pt=' '.join(vals)
  elif vals:pt=vals[0]
  rows.append(dict.fromkeys(FIELDS,'')|{'id':f'BOR-CORBO-PAU-{n+1:03d}','source':src,'segmentation':seg,'gloss':gl,'portuguese':pt})
 return rows

def main():
 if len(sys.argv)!=3:raise SystemExit('uso: import_new_interlinear.py BIBLIA.docx PAULINHO.docx')
 root=Path(__file__).resolve().parents[1]
 write(root/'CorBo_vNext/texts/biblia-oral/biblia_oral_interlinear.tsv',bible(sys.argv[1]))
 write(root/'CorBo_vNext/texts/paulinho/paulinho_interlinear.tsv',paulinho(sys.argv[2]))
 print('Importados: 20 unidades Bíblia oral; 30 unidades Texto Paulinho.')
if __name__=='__main__':main()
